#!/usr/bin/env python3
import argparse
import json
import os
import re
import sys
import urllib.error
import urllib.request
from pathlib import Path

from docx import Document
from docx.shared import Pt


DEFAULT_BASE_URL = "https://api.deepseek.com"
DEFAULT_MODEL = "deepseek-v4-pro"

CORE_RULES = """你处理的是专家访谈纪要。核心原则是：先保全信息，再重建逻辑，最后才压缩表达。

逐字稿是事实底座。访谈提纲只帮助理解研究目的，不能决定纪要结构，也不能为未回答的问题补答案。用户速记只提示重点、时间口径和可能的术语问题，不能覆盖逐字稿；速记更简略时必须保留逐字稿的完整限定条件。

不得补全、推断或发明来源之外的信息。不得归一化口头数字。不同时间、对象、状态、机制、策略、条件、例外、例子、排序或优先级不得为了简洁而合并。并列原因不得改写成因果链。实际、目标、计划、申请中、尚未批准和专家个人预计必须区分。"""

EXTRACT_SYSTEM = CORE_RULES + """

你当前只执行 Step 1：提取有效信息。不要写纪要，不要追求简洁。宁可暂时重复，也不要提前总结丢信息。完整扫描逐字稿，在主题、时间、对象、状态、机制或子问题变化时拆分信息条目，并保留足够上下文来消解“这个、它、也是一样”等指代。

输出 Markdown，格式如下：
# 有效信息清单
## 主题名称
### I001
- 来源：时间戳或最稳定的逐字稿位置
- 主题：
- 时间：未说明时写“未说明”
- 对象：未说明时写“未说明”
- 状态：已发生 / 当前实际 / 当前目标 / 计划 / 申请中 / 尚未批准 / 专家个人预计 / 未说明
- 内容：保留原始数字措辞、完整限定条件、机制和业务含义
- 关系：结论 / 原因 / 结果 / 并列原因 / 条件 / 例外 / 对比 / 举例 / 未说明

规则：
1. 同一句中若包含不同时间、对象、状态、机制或关系，拆成多条。
2. 所有有研究价值的数字、时间、对象、原因、机制、结果、措施、条件、例外、业务例子、对比、排序、优先级和不确定性都要进入清单。
3. 后文修正前文时，既记录最终口径，也注明发生过修正；不要把冲突静默抹平。
4. 访谈提纲或速记强调但逐字稿没有支持的内容，不进入事实清单；可在末尾“待核对提示”中说明未找到支持。
5. 不要输出摘要、文章或最终结论。"""

STRUCTURE_SYSTEM = CORE_RULES + """

你当前只执行 Step 2：根据完整信息清单重建专家真正讲出的业务逻辑。不要写正式纪要。不要沿用访谈问题、提纲或 Q1/Q2/Q3 的顺序。

跨问题内容可以归到同一主题，但只有时间、对象、状态和含义均一致时才能合并。相同主题下的不同时间、对象、状态或策略必须保留为不同分支。明确指出结果、原因/机制、应对措施、资源变化、未来计划之间的真实关系；并列就写并列，不要创造因果。

输出 Markdown：
# 逻辑结构
## 一、主题标题
- 主判断：...
- 逻辑分支：...
- 必须保留的信息条目：I001、I002...
- 不得合并的差异：...
- 推荐表达顺序：...

最后增加：
## 口径与冲突处理
- 列出时间、对象、状态、术语或前后修正中需要在正式纪要中明确区分的内容。

大纲必须覆盖所有有独立业务含义的信息条目；不能只挑最显眼的结论。"""

DRAFT_SYSTEM = CORE_RULES + """

你当前执行 Step 3：根据有效信息清单和逻辑结构写正式中文纪要。此时才追求 boss-facing、concise、scannable，但简洁只能来自删除口头语和无意义重复、合并真正完全重复的信息、改善句法与层级；不得删除独立业务信息。

默认用“标题 + 主结论 + 支撑 bullet”：
- **主结论：** 一句话表达核心判断。
  - 下层 bullet 保留必要的原因、机制、数字、时间、对象、策略、条件、例外、例子和出处。

不要强制 Q&A，不要每句话机械添加【事实】【观点】【待确认】，不要强制大量表格。仅在横向比较、数字对比或多对象对照明显更清楚时使用表格。用自然限定语区分“截至某时实际”“当前目标”“计划”“尚未批准”“专家个人预计”。

输出完整 Markdown 纪要，建议包含：
# 会议纪要
## 会议信息
## 核心结论
## 按专家实际逻辑组织的主题章节
## 风险、不确定性与待确认事项（仅在确有内容时）

每个具体数字和关键事实应在对应 bullet 末尾保留来源时间戳或信息条目编号。不要输出写作说明或 QA 报告。"""

AUDIT_SYSTEM = CORE_RULES + """

你当前执行 Step 4：Coverage Audit。将纪要初稿与完整信息清单、逻辑结构逐项对照，发现问题后直接修复纪要。最终只输出修正后的完整 Markdown 纪要，不输出 QA 报告、检查清单、修改说明或思考过程。

必须逐项检查并修复：
1. 重要数字是否遗漏；
2. 重要时间口径是否遗漏或混淆；
3. 不同对象是否被错误合并；
4. 并列关系是否被错误写成因果；
5. 实际、目标、计划、申请中、尚未批准、专家预测是否混淆；
6. 明确排序或优先级是否保留；
7. 重要策略的具体措施是否被过度抽象；
8. 有业务含义的例子是否被删除；
9. 后文修正/澄清是否采用最终明确口径；
10. 信息清单中由用户速记提醒且获逐字稿支持的内容是否体现；
11. 是否为了简洁删除了任何独立信息。

验收问题：用户是否还需要重新通读逐字稿才能发现重要遗漏或被改变的业务差异？如果是，继续修复后再输出。"""


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Generate high-fidelity meeting minutes through a four-stage DeepSeek workflow."
    )
    parser.add_argument("--transcript", required=True, help="Transcript Markdown path")
    parser.add_argument("--minutes-md", required=True, help="Output final minutes Markdown path")
    parser.add_argument("--minutes-docx", required=True, help="Output final minutes DOCX path")
    parser.add_argument("--topic", default="", help="Meeting topic")
    parser.add_argument("--date", default="", help="Meeting date")
    parser.add_argument("--participants", default="", help="Participants/roles")
    parser.add_argument("--outline", help="Optional interview outline file")
    parser.add_argument("--notes", help="Optional user interview notes file")
    parser.add_argument("--information-md", help="Output Step 1 information inventory")
    parser.add_argument("--structure-md", help="Output Step 2 logic structure")
    parser.add_argument("--draft-md", help="Output Step 3 draft minutes")
    parser.add_argument("--usage-json", help="Output model usage JSON")
    parser.add_argument("--prompt-dir", help="Dry-run prompt package directory")
    parser.add_argument("--dry-run", action="store_true", help="Write four-stage prompt package; do not call API")
    args = parser.parse_args()

    transcript = read_text(args.transcript)
    outline = read_optional(args.outline)
    notes = read_optional(args.notes)
    paths = artifact_paths(args)
    meta = build_meta(args)

    if args.dry_run:
        write_dry_run_prompts(paths["prompt_dir"], transcript, outline, notes, meta)
        print("PROMPT_DIR=" + str(paths["prompt_dir"]))
        return 0

    usage_records = []

    information, usage = call_deepseek(
        "step1_extract",
        EXTRACT_SYSTEM,
        build_extract_prompt(transcript, outline, notes, meta),
        temperature=0.0,
    )
    write_text(paths["information"], ensure_text(information, "Step 1"))
    usage_records.append(usage)

    structure, usage = call_deepseek(
        "step2_structure",
        STRUCTURE_SYSTEM,
        build_structure_prompt(information, meta),
        temperature=0.1,
    )
    write_text(paths["structure"], ensure_text(structure, "Step 2"))
    usage_records.append(usage)

    draft, usage = call_deepseek(
        "step3_draft",
        DRAFT_SYSTEM,
        build_draft_prompt(information, structure, meta),
        temperature=0.2,
    )
    write_text(paths["draft"], ensure_text(draft, "Step 3"))
    usage_records.append(usage)

    final_minutes, usage = call_deepseek(
        "step4_audit_repair",
        AUDIT_SYSTEM,
        build_audit_prompt(information, structure, draft, meta),
        temperature=0.0,
    )
    final_minutes = ensure_text(final_minutes, "Step 4").strip() + "\n"
    usage_records.append(usage)

    write_text(args.minutes_md, final_minutes)
    markdown_to_docx(final_minutes, args.minutes_docx)
    write_usage(paths["usage"], usage_records)

    print("INFORMATION_MD=" + str(paths["information"]))
    print("STRUCTURE_MD=" + str(paths["structure"]))
    print("DRAFT_MD=" + str(paths["draft"]))
    print("MINUTES_MD=" + str(Path(args.minutes_md)))
    print("MINUTES_DOCX=" + str(Path(args.minutes_docx)))
    print("USAGE_JSON=" + str(paths["usage"]))
    return 0


def artifact_paths(args) -> dict:
    minutes_md = Path(args.minutes_md)
    stem = minutes_md.stem
    base = stem[:-5] if stem.endswith("_会议纪要") else stem
    parent = minutes_md.parent
    return {
        "information": Path(args.information_md) if args.information_md else parent / f"{base}_信息清单.md",
        "structure": Path(args.structure_md) if args.structure_md else parent / f"{base}_逻辑结构.md",
        "draft": Path(args.draft_md) if args.draft_md else parent / f"{base}_纪要初稿.md",
        "usage": Path(args.usage_json) if args.usage_json else parent / f"{base}_模型用量.json",
        "prompt_dir": Path(args.prompt_dir) if args.prompt_dir else parent / f"{base}_dry_run_prompts",
    }


def build_meta(args) -> str:
    rows = []
    if args.topic:
        rows.append(f"会议主题：{args.topic}")
    if args.date:
        rows.append(f"会议日期：{args.date}")
    if args.participants:
        rows.append(f"参会角色：{args.participants}")
    return "\n".join(rows) if rows else "未提供"


def build_extract_prompt(transcript: str, outline: str, notes: str, meta: str) -> str:
    return f"""请从以下材料提取完整有效信息清单。

<meeting_meta>
{meta}
</meeting_meta>

<interview_outline optional="true" role="context_only_not_structure">
{outline or "未提供"}
</interview_outline>

<user_notes optional="true" role="emphasis_only_cannot_override_transcript">
{notes or "未提供"}
</user_notes>

<transcript role="fact_base">
{transcript}
</transcript>
"""


def build_structure_prompt(information: str, meta: str) -> str:
    return f"""请根据有效信息清单重建专家真正的业务逻辑结构。

<meeting_meta>
{meta}
</meeting_meta>

<information_inventory>
{information}
</information_inventory>
"""


def build_draft_prompt(information: str, structure: str, meta: str) -> str:
    return f"""请生成正式会议纪要初稿。

<meeting_meta>
{meta}
</meeting_meta>

<information_inventory>
{information}
</information_inventory>

<logic_structure>
{structure}
</logic_structure>
"""


def build_audit_prompt(information: str, structure: str, draft: str, meta: str) -> str:
    return f"""请执行 Coverage Audit，自动修复后只输出最终会议纪要。

<meeting_meta>
{meta}
</meeting_meta>

<information_inventory>
{information}
</information_inventory>

<logic_structure>
{structure}
</logic_structure>

<draft_minutes>
{draft}
</draft_minutes>
"""


def write_dry_run_prompts(prompt_dir: Path, transcript: str, outline: str, notes: str, meta: str) -> None:
    prompt_dir.mkdir(parents=True, exist_ok=True)
    prompts = {
        "01_extract.md": EXTRACT_SYSTEM + "\n\n---\n\n" + build_extract_prompt(transcript, outline, notes, meta),
        "02_structure.md": STRUCTURE_SYSTEM + "\n\n---\n\n" + build_structure_prompt("{{STEP1_INFORMATION_INVENTORY}}", meta),
        "03_draft.md": DRAFT_SYSTEM + "\n\n---\n\n" + build_draft_prompt(
            "{{STEP1_INFORMATION_INVENTORY}}", "{{STEP2_LOGIC_STRUCTURE}}", meta
        ),
        "04_audit_repair.md": AUDIT_SYSTEM + "\n\n---\n\n" + build_audit_prompt(
            "{{STEP1_INFORMATION_INVENTORY}}",
            "{{STEP2_LOGIC_STRUCTURE}}",
            "{{STEP3_DRAFT_MINUTES}}",
            meta,
        ),
    }
    for name, content in prompts.items():
        write_text(prompt_dir / name, content.strip() + "\n")


def call_deepseek(stage: str, system_prompt: str, user_prompt: str, temperature: float):
    api_key = os.environ.get("DEEPSEEK_API_KEY")
    if not api_key:
        raise SystemExit("Missing DEEPSEEK_API_KEY")
    base_url = os.environ.get("DEEPSEEK_BASE_URL", DEFAULT_BASE_URL).rstrip("/")
    model = os.environ.get("DEEPSEEK_MODEL", DEFAULT_MODEL)
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": temperature,
    }
    if os.environ.get("DEEPSEEK_MAX_TOKENS"):
        payload["max_tokens"] = int(os.environ["DEEPSEEK_MAX_TOKENS"])
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(
        f"{base_url}/chat/completions",
        data=body,
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=900) as resp:
            result = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        err = e.read().decode("utf-8", errors="replace")
        raise SystemExit(f"DeepSeek HTTP {e.code} during {stage}: {redact(err)}")
    except Exception as e:
        raise SystemExit(f"DeepSeek request failed during {stage}: {e}")
    try:
        choice = result["choices"][0]
        content = choice["message"]["content"]
    except (KeyError, IndexError, TypeError):
        raise SystemExit(f"DeepSeek response during {stage} lacked choices[0].message.content")
    finish_reason = choice.get("finish_reason")
    if finish_reason == "length":
        raise SystemExit(
            f"DeepSeek output was truncated during {stage}; no partial minutes were delivered. "
            "Increase DEEPSEEK_MAX_TOKENS or use a model with a larger output limit."
        )
    usage = {
        "stage": stage,
        "model": result.get("model", model),
        "finish_reason": finish_reason,
        **(result.get("usage") or {}),
    }
    return content, usage


def write_usage(path: Path, records) -> None:
    total = {}
    for record in records:
        for key, value in record.items():
            if key.endswith("_tokens") and isinstance(value, int):
                total[key] = total.get(key, 0) + value
    write_text(path, json.dumps({"calls": records, "total": total}, ensure_ascii=False, indent=2) + "\n")


def markdown_to_docx(markdown: str, path: str) -> None:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            doc.add_heading(clean_inline(stripped[level:].strip()), level=min(level, 3))
            i += 1
            continue
        if is_table_line(stripped) and i + 1 < len(lines) and is_separator_line(lines[i + 1]):
            table_lines = [stripped]
            i += 2
            while i < len(lines) and is_table_line(lines[i].strip()):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            indent = len(line) - len(line.lstrip())
            style = "List Bullet 2" if indent >= 2 and "List Bullet 2" in styles else "List Bullet"
            doc.add_paragraph(clean_inline(re.sub(r"^\s*[-*]\s+", "", line)), style=style)
            i += 1
            continue
        if re.match(r"^\s*\d+[.)]\s+", line):
            doc.add_paragraph(clean_inline(re.sub(r"^\s*\d+[.)]\s+", "", line)), style="List Number")
            i += 1
            continue
        doc.add_paragraph(clean_inline(stripped))
        i += 1

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def add_table(doc: Document, table_lines):
    rows = [split_table_row(line) for line in table_lines]
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=max(len(row) for row in rows))
    table.style = "Table Grid"
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            table.cell(row_idx, col_idx).text = clean_inline(value)


def is_table_line(line: str) -> bool:
    return line.startswith("|") and line.endswith("|")


def is_separator_line(line: str) -> bool:
    return bool(re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", line)) and "---" in line


def split_table_row(line: str):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def clean_inline(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    return text


def read_optional(path: str) -> str:
    return read_text(path) if path else ""


def read_text(path) -> str:
    return Path(path).read_text(encoding="utf-8")


def write_text(path, text: str) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def ensure_text(text: str, stage: str) -> str:
    if not isinstance(text, str) or not text.strip():
        raise SystemExit(f"{stage} returned empty content")
    return text.strip() + "\n"


def redact(text: str) -> str:
    return re.sub(r"(sk-[A-Za-z0-9_\-]{8})[A-Za-z0-9_\-]+", r"\1[REDACTED]", text)


if __name__ == "__main__":
    raise SystemExit(main())
